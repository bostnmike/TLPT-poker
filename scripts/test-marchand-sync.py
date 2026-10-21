"""Read-only importer regression tests; the canonical workbook is never written."""
import importlib.util
import json
from copy import deepcopy
from pathlib import Path
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from openpyxl.utils.cell import get_column_letter

root = Path(__file__).resolve().parent.parent
spec = importlib.util.spec_from_file_location("sync", root / "scripts/sync-marchand-data.py")
sync = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sync)
payload = json.loads((root / "data/marchand-pucks.json").read_text())
workbook = Workbook()
workbook.remove(workbook.active)
for name in sync.SOURCE_SHEETS:
    workbook.create_sheet(name)
for record in payload["records"]:
    sheet = workbook[record["sourceSheet"]]
    for column, field in enumerate(record["sourceData"], 1):
        sheet.cell(1, column, field["label"])
        cell = sheet.cell(record["sourceRow"], column, field["value"])
        if field.get("url"):
            cell.hyperlink = field["url"]
for name in sync.SOURCE_SHEETS:
    sheet = workbook[name]
    sheet.add_table(Table(displayName=sync.SOURCE_TABLES[name], ref=f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"))
for record in payload["records"]:
    rebuilt = deepcopy(record)
    sync.synchronize_record(rebuilt, workbook, workbook)
    assert rebuilt == record, (record["inventoryId"], "round-trip drift")
reordered = deepcopy(payload)
for record in reordered["records"]:
    record["sourceRow"] = 999
sync.discover_source_records(reordered, workbook)
assert reordered == payload, "row sorting changed artifact identities"
# Non-record tabs, stray populated rows and cells outside named tables are ignored.
workbook.create_sheet("New Puck Intake", 0).append(["ID", "Date", "Category"])
workbook.create_sheet("Dashboard Data").append([1, "fake duplicate", 172])
workbook.create_sheet("Data Quality").append([1, "fake duplicate", 172])
workbook.create_sheet("Dashboard").append([172])
stray = workbook["Goals & Games"]
for col in range(1, stray.max_column + 1):
    stray.cell(300, col, stray.cell(2, col).value)
stray.cell(2, 100, "not a source column")
again = deepcopy(payload);sync.discover_source_records(again, workbook)
assert again == payload
rebuilt = deepcopy(payload['records'][0]);sync.synchronize_record(rebuilt, workbook, workbook)
assert rebuilt == payload['records'][0], "stray cells leaked into sourceData"
assert sync.iso_date(45951) == "2025-10-21"
try:
    sync.iso_date(45951.5)
except ValueError:
    pass
else:
    raise AssertionError("Date fractions must not silently change a calendar date")
broken = deepcopy(workbook);del broken['Goals & Games'].tables[sync.SOURCE_TABLES['Goals & Games']]
try:
    sync.discover_source_records(deepcopy(payload), broken)
except ValueError:
    pass
else:
    raise AssertionError('Missing named table must fail, never fall back to used range')
by_id = {record["inventoryId"]: record for record in payload["records"]}
assert by_id[324]["category"] == "Milestone"
assert by_id[324]["period"] == "2" and by_id[324]["time"] == "4:21"
assert by_id[324]["goalType"] == "ESG / GWG"
assert by_id[324]["goalieScoredAgainst"] == "Jeremy Swayman"
assert "Palmieri" in by_id[324]["description"]
assert by_id[9]["description"] and by_id[500]["playerCodes"] == ["BM63"]
for number in (10, 12, 32, 67, 59):
    assert by_id[number]["authenticationType"] == "Digital COA / Hologram"
    assert by_id[number]["authenticationEvidence"]
print(f"PASS: all {len(by_id)} artifacts round-trip across six sheets; stable IDs, scoring details, titles, video labels and holograms preserved.")

recap = "https://www.nhl.com/news/new-jersey-devils-boston-bruins-game-recap-295228280"
assert sync.video_parts(recap, {}) == ("article", "")
reviewed = {"videoClipId": "6335868297112", "videoClipSourceUrl": recap}
assert sync.video_parts(recap, reviewed) == ("nhl", "6335868297112")
assert sync.video_parts(recap + "-changed", reviewed) == ("article", ""), "Reviewed clip must belong to this exact canonical source"

assert by_id[49]["date"] == "2022-05-06"
assert by_id[49]["authenticationSerial"] == "MGA37776"
assert by_id[49]["careerStat"] == 46 and by_id[49]["category"] == "PO Goal"
assert by_id[342]["sourceSheet"] == "Milestones" and by_id[342]["sourceRow"] == 44
assert by_id[342]["authenticationSerial"] == "XJ05532097"
assert by_id[322]["authenticationSerial"] == "XJ05532094"

assert by_id[57]["authenticationSerial"] == "MGA21933"
survivor = by_id[313]
assert survivor["sourceSheet"] == "Milestones" and survivor["sourceRow"] == 15
assert survivor["seasonStat"] == 4 and survivor["time"] == "12:49"
assert survivor["primaryAssist"] == "BM63" and survivor["secondaryAssist"] == ""
assert survivor["goalieScoredAgainst"] == "Connor Ingram"
source_video = next(field for field in survivor["sourceData"] if field["label"] == "Video URL")
assert source_video["value"] == "Watch Goal" and source_video["url"] == survivor["videoUrl"]
fresh = deepcopy(survivor)
fresh.update(videoUrl="", videoProvider="", videoId="")
sync.synchronize_record(fresh, workbook, workbook)
assert fresh == survivor, "Native hyperlink target must import without an existing website URL"

# New native hyperlink labels remain readable; raw URL cells keep the existing default.
unlabeled = deepcopy(survivor)
unlabeled["videoLabel"] = ""
sync.synchronize_record(unlabeled, workbook, workbook)
assert unlabeled["videoLabel"] == "Watch Goal"
video_column = next(column for column in range(1, workbook["Milestones"].max_column + 1) if workbook["Milestones"].cell(1, column).value == "Video URL")
workbook["Milestones"].cell(survivor["sourceRow"], video_column).value = survivor["videoUrl"]
unlabeled["videoLabel"] = ""
sync.synchronize_record(unlabeled, workbook, workbook)
assert unlabeled["videoLabel"] == "Game Context"

assert by_id[59]["authenticationSerial"] == "G009264"
renumbered=by_id[300]
assert renumbered["sourceSheet"]=="Milestones" and renumbered["sourceRow"]==2
assert renumbered["category"]=="RS Point" and renumbered["puckType"]=="Goal Scored Puck"
assert renumbered["scorerCode"]=="EL27" and renumbered["primaryAssist"]=="BM63" and renumbered["secondaryAssist"]=="AE5"
assert renumbered["authenticationSerial"]=="BB25087"
assert renumbered["sourceData"][22]["value"]=="images/pucks/59/front.webp"

assert sync.video_parts("https://www.nhl.com/video/c-6335822476112", {}) == ("nhl", "6335822476112")
for invalid in ["https://www.nhl.com/gamecenter/2021030133", "https://www.nhl.com/video/c-2021030133"]:
    try:
        sync.video_parts(invalid, {})
    except ValueError:
        pass
    else:
        raise AssertionError("A game ID must not become a video clip: " + invalid)
