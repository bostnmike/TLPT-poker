#!/usr/bin/env python3

import argparse
import hashlib
import json
import re
from datetime import date, datetime
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.datetime import from_excel


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DATA = ROOT / "data/marchand-pucks.json"
CODE_PATTERN = re.compile(r"\b[A-Za-z]{1,3}\d{1,2}\b")
SOURCE_SHEETS = ("Goals & Games", "Milestones", "Road to History", "Road to Repeat", "Hockey Fights Cancer", "Warm-Up Pucks")
DATABASE_SCHEMA = json.loads((ROOT / "docs/marchand-database-schema-20261020.json").read_text())
SOURCE_TABLES = {sheet: table["name"] for sheet, table in DATABASE_SCHEMA["record_tables"].items()}
NORMALIZED_FIELDS = {
    "Career Stat Basis": "careerStatBasis", "Season Stat Basis": "seasonStatBasis",
    "Goal Strength": "goalStrength", "Recorded Scorer": "recordedScorer",
    "NHL Game ID": "nhlGameId", "NHL Goal Event ID": "nhlGoalEventId",
}
ARENA_ALIASES = {row[0]: row[1] for catalog in DATABASE_SCHEMA["catalogs"] if catalog["name"] == "ArenaAliases" for row in catalog["rows"]}
VENUES = {row[0]: row[1] for catalog in DATABASE_SCHEMA["catalogs"] if catalog["name"] == "PhysicalVenues" for row in catalog["rows"]}
SOURCE_ID_HEADERS = {
    "Goals & Games": "ID",
    "Milestones": "ID #",
    "Road to History": "Inventory ID",
    "Road to Repeat": "Inventory ID",
    "Hockey Fights Cancer": "ID #",
    "Warm-Up Pucks": "ID #",
}
KEY_PREFIXES = {
    "Goals & Games": "goal",
    "Milestones": "milestone",
    "Road to History": "history",
    "Road to Repeat": "repeat",
    "Hockey Fights Cancer": "hfc",
    "Warm-Up Pucks": "warmup",
}


def parse_args():
    parser = argparse.ArgumentParser(
        description="Synchronize the Marchand site database from the canonical workbook."
    )
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--data", type=Path, default=DEFAULT_DATA)
    return parser.parse_args()


def blank_to_empty(value):
    return "" if value is None else value


def as_number(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        return int(value)
    if isinstance(value, str) and re.fullmatch(r"-?\d+", value.strip()):
        return int(value)
    return value


def iso_date(value, epoch=None):
    if value in (None, ""):
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value != int(value):
            raise ValueError(f"Calendar date has a time fraction: {value!r}")
        return from_excel(value, **({"epoch": epoch} if epoch else {})).strftime("%Y-%m-%d")
    text = str(value).strip()
    for pattern in ("%B %d, %Y", "%B %-d, %Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).strftime("%Y-%m-%d")
        except (ValueError, TypeError):
            continue
    raise ValueError(f"Unsupported date value: {value!r}")


def source_text(header, value, epoch=None):
    if header == "Date":
        return iso_date(value, epoch)
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def player_codes(*values):
    result = []
    for value in values:
        for code in CODE_PATTERN.findall(str(value or "")):
            if code.lower() == "ot":
                continue
            if code not in result:
                result.append(code)
    return result


def video_parts(url, previous):
    if not url:
        return "", ""
    if (
        url == previous.get("videoUrl")
        and previous.get("videoProvider")
        and previous.get("videoId")
        and "nhl.com" not in urlparse(url).netloc.lower()
    ):
        return previous.get("videoProvider", ""), previous.get("videoId", "")
    parsed = urlparse(url)
    host = parsed.netloc.lower()
    if "youtube.com" in host:
        return "youtube", parse_qs(parsed.query).get("v", [""])[0]
    if "youtu.be" in host:
        return "youtube", parsed.path.strip("/").split("/")[0]
    if "nhl.com" in host:
        if parsed.path.startswith("/news/"):
            # A recap article ID is not a Brightcove video ID. A reviewed clip
            # may accompany the canonical article; otherwise retain the recap only.
            clip_id = str(previous.get("videoClipId", ""))
            if clip_id and previous.get("videoClipSourceUrl") == url:
                if not re.fullmatch(r"\d{13}", clip_id):
                    raise ValueError(f"Invalid reviewed NHL clip ID: {clip_id}")
                return "nhl", clip_id
            return "article", ""
        if "/video/" not in parsed.path:
            raise ValueError(f"NHL game summaries are not video clip URLs: {url}")
        match = re.search(r"(?<!\d)(\d{13})$", parsed.path.rstrip("/"))
        if not match:
            raise ValueError(f"NHL video URL has no valid 13-digit clip ID: {url}")
        return "nhl", match.group(1)
    raise ValueError(f"Unsupported video URL: {url}")


def source_table(sheet):
    """Only explicitly named collection tables own records and source columns."""
    name = SOURCE_TABLES[sheet.title]
    if name not in sheet.tables:
        raise ValueError(f"Missing source table {name!r} in {sheet.title!r}")
    bounds = range_boundaries(sheet.tables[name].ref)
    left, header_row, right, bottom = bounds
    headers = [sheet.cell(header_row, col).value for col in range(left, right + 1)]
    required = DATABASE_SCHEMA["record_tables"][sheet.title]["headers"]
    if len(headers) != len(set(headers)) or any(not h for h in headers):
        raise ValueError(f"Blank or duplicate headers in {name}")
    missing = set(required) - set(headers)
    if missing:
        raise ValueError(f"Missing known headers in {name}: {sorted(missing)}")
    return bounds, {h: col for col, h in enumerate(headers, left)}


def source_data(sheet_values, sheet_formulas, row):
    fields = []
    (left, header_row, right, bottom), _ = source_table(sheet_values)
    if not header_row < row <= bottom:
        raise ValueError(f"Source row outside named table: {sheet_values.title}!{row}")
    for column in range(left, right + 1):
        header = sheet_values.cell(header_row, column).value
        value = sheet_values.cell(row, column).value
        formula_cell = sheet_formulas.cell(row, column)
        url = formula_cell.hyperlink.target if formula_cell.hyperlink else ""
        fields.append({"label": header, "value": source_text(header, value, sheet_values.parent.epoch), "url": url})
    return fields


def field_map(fields):
    return {field["label"]: field["value"] for field in fields}


def synchronize_normalized_fields(record, values):
    for label, key in NORMALIZED_FIELDS.items():
        record[key] = values.get(label, "")
    for label, key in (("Game-Winning Goal", "gameWinningGoal"), ("Overtime Goal", "overtimeGoal")):
        value = values.get(label, "")
        if value not in ("", "0", "1"):
            raise ValueError(f"Invalid scoring flag {label}: {value!r}")
        record[key] = None if value == "" else value == "1"
    venue_key = ARENA_ALIASES.get(record["arena"])
    if not venue_key:
        raise ValueError(f"Unmapped arena label: {record['arena']!r}")
    record["venueKey"], record["venueName"] = venue_key, VENUES[venue_key]


def synchronize_record(record, workbook_values, workbook_formulas):
    sheet_name = record["sourceSheet"]
    row = record["sourceRow"]
    values_sheet = workbook_values[sheet_name]
    formulas_sheet = workbook_formulas[sheet_name]
    fields = source_data(values_sheet, formulas_sheet, row)
    values = field_map(fields)

    record["sourceData"] = fields
    if sheet_name == "Goals & Games":
        record.update(
            inventoryId=int(values["ID"]),
            homeRoad=values["Home/Road"],
            team=values["Marchand Team"],
            category=values["Category"],
            careerStat=as_number(values["Career Stat"]),
            seasonStat=as_number(values["Season Stat"]),
            description=values.get("Description", record.get("description", "")),
            date=values["Date"],
            arena=values["Arena"],
            opponent=values["Opponent"],
            period=values["Period"],
            time=values["Time"],
            goalType=values["Goal Type"],
            primaryAssist=values["Primary Assist"],
            secondaryAssist=values["Secondary Assist"],
            score=values["Score"],
            puckType=values["Puck Type"],
            notes=values["Notes"],
            goalieScoredAgainst=values["Goalie Scored Against"],
            playerCodes=player_codes(values["Primary Assist"], values["Secondary Assist"]),
        )
        canonical_url = (values["Goal Video"] if values["Goal Video"].startswith("https://")
                         else next(field["url"] for field in fields if field["label"] == "Goal Video"))
        canonical_label = values["Goal Video"]
        record["videoLabel"] = canonical_label
        record["videoUrl"] = canonical_url
        record["videoProvider"], record["videoId"] = video_parts(canonical_url, record)
    elif sheet_name in ("Milestones", "Hockey Fights Cancer", "Warm-Up Pucks"):
        record.update(
            inventoryId=int(values["ID #"]),
            homeRoad=values.get("Home/Road", ""),
            team=values["Marchand Team"],
            category=values["Category"],
            careerStat=as_number(values.get("Career Stat")),
            seasonStat=as_number(values.get("Season Stat")),
            description=values["Description"],
            date=values["Date"],
            arena=values["Arena"],
            opponent=values["Opponent"],
            period=values.get("Period", ""),
            time=values.get("Time", ""),
            goalType=values.get("Goal Type", ""),
            primaryAssist=values.get("Primary Assist", ""),
            secondaryAssist=values.get("Secondary Assist", ""),
            score=values.get("Score", ""),
            puckType=values["Puck Type"],
            notes=values["Notes"],
            game=None,
            wins=None,
            losses=None,
            points=None,
            goalieScoredAgainst=values.get("Goalie Scored Against", ""),
            playerCodes=record.get("playerCodes") or player_codes(values["Description"], values["Notes"]),
        )
    elif sheet_name == "Road to History":
        wins = as_number(values["W"])
        losses = as_number(values["L"])
        overtime_losses = as_number(values["OTL"])
        record.update(
            inventoryId=int(values["Inventory ID"]),
            homeRoad=values["Home/Road"],
            team=values["Marchand Team"],
            category=values["Category"],
            careerStat=None,
            seasonStat=None,
            description=values["Description"],
            date=values["Date"],
            arena=values["Arena"],
            opponent=values["Opponent"],
            period=values.get("Period", ""),
            time=values.get("Time", ""),
            goalType=values.get("Goal Type", ""),
            primaryAssist=values.get("Primary Assist", ""),
            secondaryAssist=values.get("Secondary Assist", ""),
            score=values["Score"],
            puckType=values["Puck Type"],
            notes=values["Notes"],
            game=as_number(values["Game"]),
            wins=wins,
            losses=losses,
            overtimeLosses=overtime_losses,
            seasonRecord=f"{wins}-{losses}-{overtime_losses}",
            points=as_number(values["Pts"]),
            goalieScoredAgainst=values.get("Goalie Scored Against", ""),
            playerCodes=record.get("playerCodes") or player_codes(values["Description"], values["Notes"]),
        )
    elif sheet_name == "Road to Repeat":
        wins, losses = as_number(values["Playoff Wins"]), as_number(values["Playoff Losses"])
        record.update(
            inventoryId=int(values["Inventory ID"]), homeRoad=values["Home/Road"],
            team=values["Marchand Team"], category=values["Category"],
            description=values["Description"], date=values["Date"], arena=values["Arena"],
            opponent=values["Opponent"], score=values["Score"], puckType=values["Puck Type"],
            notes=values["Notes"], game=as_number(values["Game"]),
            playoffRound=values["Playoff Round"], seriesRecord=values["Series Record"],
            playoffRecord=f"{wins}–{losses}" if wins is not None and losses is not None else "",
            careerStat=as_number(values.get("Career Stat")),
            seasonStat=as_number(values.get("Season Stat")),
            period=values.get("Period", ""), time=values.get("Time", ""),
            goalType=values.get("Goal Type", ""),
            primaryAssist=values.get("Primary Assist", ""),
            secondaryAssist=values.get("Secondary Assist", ""),
            goalieScoredAgainst=values.get("Goalie Scored Against", ""),
            wins=None, losses=None, points=None, playerCodes=["BM63"],
        )
        record.pop("scorerCode", None)
        if record["category"] == "Assist":
            scorer = re.search(r"\bon ([A-Za-z]{1,3}\d{1,2}) goal\b", record["notes"], re.I)
            if scorer:
                record["scorerCode"] = scorer.group(1)
                record["playerCodes"] = player_codes(
                    record["scorerCode"], record["primaryAssist"], record["secondaryAssist"]
                )
    else:
        raise ValueError(f"Unsupported source sheet: {sheet_name}")

    if sheet_name != "Goals & Games" and values.get("Video URL"):
        video_field = next(field for field in fields if field["label"] == "Video URL")
        canonical_url = (values["Video URL"] if values["Video URL"].startswith("https://")
                         else video_field["url"])
        provider, video_id = video_parts(canonical_url, record)
        record["videoUrl"] = canonical_url
        record["videoLabel"] = record.get("videoLabel") or (
            "Game Context" if values["Video URL"].startswith("https://") else values["Video URL"]
        )
        record["videoProvider"], record["videoId"] = provider, video_id
    if sheet_name == "Milestones" and values.get("Original Front Image Filename"):
        record["playerCodes"] = ["BM63"]
    if "Hockey Fights Cancer" in values:
        record["hockeyFightsCancer"] = values["Hockey Fights Cancer"].strip().lower() in ("yes", "true", "1")
    if "Date Status" in values:
        record["dateStatus"] = values["Date Status"]
    for label, key in (("Authentication Type", "authenticationType"),
                       ("Authentication Evidence", "authenticationEvidence"),
                       ("Authentication Notes", "authenticationNotes"),
                       ("Authentication Issuer", "authenticationIssuer"),
                       ("Hologram Number", "authenticationSerial"),
                       ("Authentication Record URL", "authenticationRecordUrl"),
                       ("Additional Authentication Issuer", "authenticationAlternateIssuer"),
                       ("Additional Hologram Number", "authenticationAlternateSerial"),
                       ("Additional Authentication Record URL", "authenticationAlternateRecordUrl"),
                       ("Authentication Lookup Instructions", "authenticationLookupInstructions")):
        if label in values:
            record[key] = values[label]
    synchronize_normalized_fields(record, values)


def discover_source_records(payload, workbook_values):
    """Match stable inventory IDs, so sorting rows cannot move an exhibit's identity."""
    def row_is_populated(sheet_name, row):
        sheet = workbook_values[sheet_name]
        (_, header_row, _, bottom), headers = source_table(sheet)
        if row <= header_row or row > bottom:
            return False
        required = (SOURCE_ID_HEADERS[sheet_name], "Marchand Team", "Category", "Date")
        populated = [sheet.cell(row, headers[header]).value not in (None, "") for header in required]
        if any(populated) and not all(populated):
            raise ValueError(f"Incomplete collection row in named table: {sheet_name}!{row}")
        return all(populated)

    canonical_rows = {}
    for sheet_name in SOURCE_SHEETS:
        sheet = workbook_values[sheet_name]
        (_, header_row, _, bottom), headers = source_table(sheet)
        for row in range(header_row + 1, bottom + 1):
            if not row_is_populated(sheet_name, row):
                continue
            value = sheet.cell(row, headers[SOURCE_ID_HEADERS[sheet_name]]).value
            if isinstance(value, bool) or not re.fullmatch(r"[1-9]\d*", str(value)):
                raise ValueError(f"Invalid inventory ID in {sheet_name}!{row}: {value!r}")
            inventory_id = int(value)
            if inventory_id in canonical_rows:
                raise ValueError(f"Duplicate canonical inventory ID: {inventory_id}")
            canonical_rows[inventory_id] = (sheet_name, row)
    payload["records"] = [record for record in payload["records"] if record["inventoryId"] in canonical_rows]
    for record in payload["records"]:
        record["sourceSheet"], record["sourceRow"] = canonical_rows[record["inventoryId"]]
    records = payload["records"]
    existing_sources = {
        (record["sourceSheet"], record["sourceRow"])
        for record in records
    }
    existing_keys = {record["key"] for record in records}

    for sheet_name in SOURCE_SHEETS:
        sheet = workbook_values[sheet_name]
        (_, header_row, _, bottom), headers = source_table(sheet)
        id_column = headers[SOURCE_ID_HEADERS[sheet_name]]
        for row in range(header_row + 1, bottom + 1):
            if not row_is_populated(sheet_name, row):
                continue
            source_key = (sheet_name, row)
            if source_key in existing_sources:
                continue
            inventory_id = as_number(sheet.cell(row, id_column).value)
            if inventory_id in (None, ""):
                continue
            record_key = f"{KEY_PREFIXES[sheet_name]}-{inventory_id}"
            if record_key in existing_keys:
                raise ValueError(
                    f"New canonical row {sheet_name}!{row} duplicates key {record_key}"
                )
            records.append(
                {
                    "key": record_key,
                    "sourceSheet": sheet_name,
                    "sourceRow": row,
                    "game": None,
                    "wins": None,
                    "losses": None,
                    "points": None,
                }
            )
            existing_sources.add(source_key)
            existing_keys.add(record_key)


def main():
    args = parse_args()
    workbook_path = args.workbook.resolve()
    data_path = args.data.resolve()
    payload = json.loads(data_path.read_text(encoding="utf-8"))

    workbook_values = load_workbook(workbook_path, data_only=True, read_only=False)
    workbook_formulas = load_workbook(workbook_path, data_only=False, read_only=False)
    discover_source_records(payload, workbook_values)
    for record in payload["records"]:
        synchronize_record(record, workbook_values, workbook_formulas)

    counts = {
        sheet: sum(record["sourceSheet"] == sheet for record in payload["records"])
        for sheet in SOURCE_SHEETS
    }
    record_count = len(payload["records"])
    payload["meta"].update(
        source=workbook_path.name,
        sourceSha256=hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        records=record_count,
        goalsAndGames=counts["Goals & Games"],
        milestones=counts["Milestones"],
        roadToHistory=counts["Road to History"],
        roadToRepeat=counts["Road to Repeat"],
        hockeyFightsCancer=counts["Hockey Fights Cancer"],
        warmUpPucks=counts["Warm-Up Pucks"],
        videos=sum(record.get("videoProvider") in ("nhl", "youtube") and bool(record.get("videoId")) for record in payload["records"]),
        sourceSheets=counts,
    )

    data_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                "data": str(data_path),
                "records": record_count,
                "videos": payload["meta"]["videos"],
                "sourceSha256": payload["meta"]["sourceSha256"],
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
